"""Local-only historical XLSX normalisation. No HTTP and no live DB imports."""
import hashlib
import json
import re
from datetime import datetime, date
from pathlib import Path
from collections import Counter

FIELDS = {
    'applicant_name': ('фио',), 'birth_date': ('дата рождения',),
    'university': ('вуз',), 'faculty': ('факультет',), 'course': ('курс',),
    'department': ('кафедра',), 'scientific_interests': ('сфера научных интересов',),
    'experience': ('опыт работы',), 'foreign_languages': ('владение иностранными языками',),
    'desired_direction': ('интересующие виды волонтёрства',),
    'expectations': ('чего вы ожидаете',), 'hobbies': ('ваше хобби',),
    'superpower': ('в чем ваша суперсила',), 'source': ('откуда вы узнали',),
    'expert_score': ('сколько баллов',), 'accepted': ('взяли ли',),
    'active': ('еще состоит',), 'activity_hours_year': ('активность в часах',),
    'tenure_years': ('длительность пребывания',), 'attendance': ('явка',),
}


def normalized(value):
    return ' '.join(str('' if value is None else value).casefold().replace('ё', 'е').split())


def unknown(v):
    return v is None or normalized(v) in {'', '-', '—', 'нет данных', 'неизвестно'}


def outcome(value):
    if unknown(value): return None
    text = normalized(value)
    if text in {'да', 'yes', 'true', 'взяли'} or value is True: return True
    if text in {'нет', 'no', 'false', 'не взяли'} or value is False: return False
    # Numeric codes are ambiguous; preserve raw, never infer rejection from 0.
    return None


def score(value):
    if unknown(value): return None
    try:
        n = float(str(value).replace(',', '.'))
        return int(n) if n in (0, 1, 2, 3) else None
    except (ValueError, TypeError): return None


def birth(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) and value.year > 1900 else value.isoformat() if value.year > 1900 else None
    text = str('' if value is None else value).strip()
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            d = datetime.strptime(text, fmt).date()
            return d.isoformat() if d.year > 1900 else None
        except ValueError: pass
    # A year alone or a serial without date formatting is not a full birth date.
    return None


def json_value(v):
    return v.isoformat() if isinstance(v, (datetime, date)) else v


def read_sources(sources):
    from openpyxl import load_workbook
    records, manifests = [], []
    for source in sources:
        path = Path(source)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            mapping = {}
            for field, aliases in FIELDS.items():
                matches = [i for i, v in enumerate(header) if any(normalized(v).startswith(normalized(a)) for a in aliases)]
                if len(matches) > 1: raise ValueError('Ambiguous historical column: ' + field)
                if matches: mapping[field] = matches[0]
            required = {'expert_score', 'accepted', 'birth_date', 'university', 'expectations'}
            if not required <= mapping.keys(): raise ValueError('Historical header mismatch: ' + sheet.title)
            count = 0
            for number, row in enumerate(rows, 2):
                if not any(v is not None and str(v).strip() for v in row): continue
                raw = {field: json_value(row[i] if i < len(row) else None) for field, i in mapping.items()}
                answers = {f: raw.get(f) for f in FIELDS if f not in {'expert_score', 'accepted', 'active', 'activity_hours_year', 'tenure_years', 'attendance'}}
                answers['birth_date'] = birth(row[mapping['birth_date']])
                fingerprint = hashlib.sha256(json.dumps({f: normalized(v) for f, v in answers.items() if f != 'applicant_name'}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                records.append({'candidate_id': hashlib.sha256(f'{digest}:{sheet.title}:{number}'.encode()).hexdigest(),
                    'source_hash': digest, 'sheet': sheet.title, 'row': number,
                    'season': '2024' if '2024' in path.name else '2025' if '2025' in path.name else 'unknown',
                    'answers': answers, 'expert_score': score(raw.get('expert_score')),
                    'accepted': outcome(raw.get('accepted')), 'active': outcome(raw.get('active')),
                    'activity_hours_year': raw.get('activity_hours_year'), 'fingerprint': fingerprint,
                    'raw': raw, 'missing_fields': [f for f in ('motivation', 'availability_hours_month', 'status', 'email', 'vk_url') if f not in mapping]})
                count += 1
            manifests.append({'source_hash': digest, 'sheet': sheet.title, 'count': count})
        workbook.close()
    counts = Counter(r['fingerprint'] for r in records)
    return records, {'sheets': manifests, 'total': len(records), 'unique_fingerprints': len(counts),
        'duplicate_rows': sum(n-1 for n in counts.values()), 'expert_labels': dict(Counter(str(r['expert_score']) for r in records)),
        'accepted_unknown': sum(r['accepted'] is None for r in records), 'network_calls': 0}
